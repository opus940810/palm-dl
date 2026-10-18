"""Convert manuscript.md into a submission-style DOCX (figures included)."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "manuscript" / "manuscript.md"
OUT = ROOT / "manuscript" / "PalmDL_manuscript.docx"
FIG = ROOT / "figures"

FIGURES = [
    ("fig1_dataset.png", "Figure 1. Dataset composition."),
    ("fig2_leakage.png", "Figure 2. Split protocol determines measured performance."),
    ("fig3_roc.png", "Figure 3. Discrimination under cluster-held-out cross-validation."),
    ("fig4_protocols.png", "Figure 4. Performance on human sites under three evaluation regimes."),
    ("fig5_calibration.png", "Figure 5. Calibration."),
    ("fig6_motif.png", "Figure 6. Sequence motif learned from the data."),
    ("fig7_species_matrix.png", "Figure 7. The palmitoylation code is eukaryotic."),
    ("fig8_interpretation.png", "Figure 8. What the model uses and how far its predictions can be trusted."),
    ("fig9_proteome_scan.png", "Figure 9. Proteome-wide scan and FDR-controlled shortlist."),
]


def add_runs(par, text: str) -> None:
    """Add inline **bold** and *italic* runs to a paragraph."""
    tokens = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            r = par.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1])
            r.italic = True
        else:
            par.add_run(tok)


def is_table_row(line: str) -> bool:
    return line.strip().startswith("|") and line.strip().endswith("|")


def add_table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(h.strip())
        r.bold = True
        r.font.size = Pt(8)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row[: len(header)]):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            add_runs(p, val.strip())
            for run in p.runs:
                run.font.size = Pt(8)
    doc.add_paragraph()


def add_result_tables(doc: Document) -> None:
    """Insert Tables 1-3 generated from the experiment outputs."""
    import pandas as pd

    tab_dir = ROOT / "results" / "tables"
    name = {"pwm": "PWM (CSS-Palm-style)", "lr": "Logistic regression", "rf": "Random forest",
            "cnn": "PalmDL-CNN", "esm": "ESM-2 view", "ensemble": "PalmDL ensemble"}

    p1 = tab_dir / "table1_cv.csv"
    if p1.exists():
        d = pd.read_csv(p1)
        doc.add_heading("Table 1. Cluster-held-out cross-validation (all species, 5 folds).", level=2)
        rows = [[name.get(r.model, r.model), f"{r.nat_auc:.3f}", f"{r.nat_aupr:.3f}",
                 f"{r.nat_mcc:.3f}", f"{r.bal_mcc:.3f}", f"{r.nat_acc:.3f}",
                 f"{r.nat_sn:.3f}", f"{r.nat_sp:.3f}"] for r in d.itertuples()]
        add_table(doc, ["Model", "AUC", "AUPR", "MCC", "Balanced MCC", "Accuracy", "Sn", "Sp"], rows)
        doc.add_paragraph("6,805 positives and 64,789 negatives (71,594 sites, 23,897 proteins). "
                          "MCC and accuracy use thresholds selected on inner validation folds; "
                          "balanced MCC uses class-balanced test subsamples.").runs[0].font.size = Pt(8)

    p2 = tab_dir / "table6_leakage.csv"
    if p2.exists():
        d = pd.read_csv(p2)
        doc.add_heading("Table 2. Leakage audit across four splitting protocols.", level=2)
        lab = {"site_random": "random site-level", "protein": "protein-level",
               "cluster": "cluster-level", "species_human": "species holdout"}
        rows = [[lab.get(r.protocol, r.protocol), str(int(r.n_train_pos)), str(int(r.n_test)),
                 f"{r.auc:.3f}", f"{r.mcc_nat:.3f}", f"{r.mcc_bal:.3f}",
                 f"{r.leaked_pos_windows * 100:.1f}%"] for r in d.itertuples()]
        add_table(doc, ["Protocol", "Training positives", "Test sites", "AUC", "MCC",
                        "Balanced MCC", "Shared windows"], rows)
        doc.add_paragraph("The same random forest and features were used throughout; only the split "
                          "changed. 'Shared windows' is the fraction of test positives whose exact "
                          "31-residue window occurs among training positives.").runs[0].font.size = Pt(8)

    p3 = tab_dir / "table2_cv_human.csv"
    if p3.exists():
        multi = pd.read_csv(p3).set_index("model")
        honly_p = tab_dir / "table3_cv_human_only.csv"
        sp_p = tab_dir / "table4_species.csv"
        honly = pd.read_csv(honly_p).set_index("model") if honly_p.exists() else None
        sp = pd.read_csv(sp_p).set_index("model") if sp_p.exists() else None
        doc.add_heading("Table 3. Human-site performance under three evaluation regimes.", level=2)
        rows = []
        for m in ["pwm", "lr", "rf", "cnn", "esm", "ensemble"]:
            if m not in multi.index:
                continue
            row = [name.get(m, m), f"{multi.loc[m, 'nat_auc']:.3f}", f"{multi.loc[m, 'bal_mcc']:.3f}"]
            row += [f"{honly.loc[m, 'nat_auc']:.3f}", f"{honly.loc[m, 'bal_mcc']:.3f}"] if honly is not None and m in honly.index else ["—", "—"]
            row += [f"{sp.loc[m, 'auc_natural']:.3f}", f"{sp.loc[m, 'mcc_balanced']:.3f}"] if sp is not None and m in sp.index else ["—", "—"]
            rows.append(row)
        add_table(doc, ["Model", "Multi-species AUC", "Multi-species bal. MCC",
                        "Human-only AUC", "Human-only bal. MCC",
                        "Species holdout AUC", "Species holdout bal. MCC"], rows)
        doc.add_paragraph("Human test sites: 572 positives in 316 proteins. The species holdout uses "
                          "no human data and removes every cluster containing a human protein.").runs[0].font.size = Pt(8)


def main() -> None:
    lines = MD.read_text(encoding="utf-8").splitlines()
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2.4)
        section.top_margin = section.bottom_margin = Cm(2.2)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.4

    i = 0
    in_refs = False
    while i < len(lines):
        line = lines[i].rstrip()
        if not line or line.strip() == "---":
            i += 1
            continue
        if line.startswith("| "):
            block = []
            while i < len(lines) and is_table_row(lines[i]):
                block.append(lines[i])
                i += 1
            rows = [[c for c in r.strip().strip("|").split("|")] for r in block]
            rows = [r for r in rows if not all(set(c.strip()) <= set("-: ") for c in r)]
            if len(rows) >= 2:
                add_table(doc, rows[0], rows[1:])
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line[2:].strip())
            r.bold = True
            r.font.size = Pt(15)
            p.paragraph_format.space_after = Pt(10)
        elif line.startswith("## "):
            h = line[3:].strip()
            in_refs = h.lower().startswith("references")
            hd = doc.add_heading(h, level=1)
            for r in hd.runs:
                r.font.color.rgb = RGBColor(0, 0, 0)
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)
        elif line.startswith("### "):
            hd2 = doc.add_heading(line[4:].strip(), level=2)
            for r in hd2.runs:
                r.font.color.rgb = RGBColor(0, 0, 0)
                r.font.name = "Times New Roman"
                r.font.size = Pt(11)
        elif line.startswith("- ") or re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number" if re.match(r"^\d+\.\s", line) else "List Bullet")
            add_runs(p, re.sub(r"^(-|\d+\.)\s+", "", line))
        elif line.startswith("**Figure") or line.startswith("**Table"):
            p = doc.add_paragraph()
            add_runs(p, line)
            for r in p.runs:
                r.font.size = Pt(9)
        else:
            p = doc.add_paragraph()
            add_runs(p, line)
            if in_refs:
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.15
                for r in p.runs:
                    r.font.size = Pt(9)
        i += 1

    add_result_tables(doc)

    doc.add_page_break()
    doc.add_heading("Figures", level=1)
    for fname, caption in FIGURES:
        path = FIG / fname
        if path.exists():
            doc.add_picture(str(path), width=Cm(16.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            p = doc.add_paragraph()
            r = p.add_run(caption)
            r.italic = True
            r.font.size = Pt(9)

    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
